import io
from typing import Dict, List, Tuple

import pandas as pd
import streamlit as st
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

APP_TITLE = "Osztálystatisztika készítő"
FOOTER_TEXT = "Készítette: Pálfi Gergő 2025"

CEL_MEZOK = [
    "Dicséret",
    "Osztályfőnöki dicséret",
    "Szaktanári dicséret",
    "Igazgatói dicséret",
    "Nevelőtestületi dicséret",
    "Házi feladat hiány",
    "Fegyelmezetlenség",
    "Késés",
    "Osztályfőnöki figyelmeztetés",
    "Szaktanári figyelmeztetés",
    "Napközis nevelői figyelmeztetés",
    "Igazgatói figyelmeztetés",
    "Osztályfőnöki intés",
    "Szaktanári intés",
    "Igazgatói intés",
    "Osztályfőnöki megrovás",
    "Szaktanári megrovás",
    "Igazgatói megrovás",
    "Na",
    "Tájékoztatás",
]

DICSERETEK = [
    "Dicséret",
    "Osztályfőnöki dicséret",
    "Szaktanári dicséret",
    "Igazgatói dicséret",
    "Nevelőtestületi dicséret",
]

FEGYELMI = [x for x in CEL_MEZOK if x not in DICSERETEK and x not in {"Na", "Tájékoztatás"}]


def parse_int(value) -> int:
    if value is None or value == "":
        return 0
    try:
        return int(float(value))
    except Exception:
        return 0


def statisztika_keszitese(xlsx_bytes: bytes, tipus: str, csak_lapok: List[str] | None = None) -> Tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    """Összesítés minden munkalapból.

    Az eredeti program logikája szerint:
    - féléves: L oszlop, Python indexben 11
    - évvégi: M oszlop, Python indexben 12
    - a kategória megnevezése az A oszlopban van
    - 10. sortól olvasunk
    """
    oszlop_index = 11 if tipus == "féléves" else 12
    wb = load_workbook(io.BytesIO(xlsx_bytes), data_only=True)
    sheetnames = csak_lapok or wb.sheetnames

    osszesito: Dict[str, int] = {kulcs: 0 for kulcs in CEL_MEZOK}
    laponkenti_rows = []

    for nev in sheetnames:
        if nev not in wb.sheetnames:
            continue
        ws = wb[nev]
        lap_osszesito: Dict[str, int] = {kulcs: 0 for kulcs in CEL_MEZOK}

        for sor in ws.iter_rows(min_row=10, max_col=oszlop_index + 1, values_only=True):
            tipus_szoveg = str(sor[0]).strip() if sor and sor[0] else ""
            if tipus_szoveg in CEL_MEZOK:
                ertek = parse_int(sor[oszlop_index])
                osszesito[tipus_szoveg] += ertek
                lap_osszesito[tipus_szoveg] += ertek

        for kategoria in CEL_MEZOK:
            laponkenti_rows.append({"Munkalap": nev, "Kategória": kategoria, "Darab": lap_osszesito[kategoria]})

    eredmeny_df = pd.DataFrame({"Kategória": CEL_MEZOK, "Darab": [osszesito[k] for k in CEL_MEZOK]})
    total = int(eredmeny_df["Darab"].sum())
    eredmeny_df["Megoszlás (%)"] = (eredmeny_df["Darab"] / total * 100).round(2) if total else 0.0

    laponkenti_df = pd.DataFrame(laponkenti_rows)
    pivot_df = laponkenti_df.pivot_table(index="Munkalap", columns="Kategória", values="Darab", aggfunc="sum", fill_value=0)
    pivot_df = pivot_df.reindex(columns=CEL_MEZOK, fill_value=0).reset_index()

    return eredmeny_df, laponkenti_df, pivot_df


def osszefoglalo_df(eredmeny_df: pd.DataFrame) -> pd.DataFrame:
    dics_sum = int(eredmeny_df.loc[eredmeny_df["Kategória"].isin(DICSERETEK), "Darab"].sum())
    fegy_sum = int(eredmeny_df.loc[eredmeny_df["Kategória"].isin(FEGYELMI), "Darab"].sum())
    taj_sum = int(eredmeny_df.loc[eredmeny_df["Kategória"].eq("Tájékoztatás"), "Darab"].sum())
    na_sum = int(eredmeny_df.loc[eredmeny_df["Kategória"].eq("Na"), "Darab"].sum())
    total = int(eredmeny_df["Darab"].sum())
    return pd.DataFrame(
        [
            {"Csoport": "Dicséretek", "Darab": dics_sum},
            {"Csoport": "Fegyelmi / figyelmeztetés jellegű", "Darab": fegy_sum},
            {"Csoport": "Tájékoztatás", "Darab": taj_sum},
            {"Csoport": "Na", "Darab": na_sum},
            {"Csoport": "Összesen", "Darab": total},
        ]
    )


def excel_export_bytes(eredmeny_df: pd.DataFrame, laponkenti_df: pd.DataFrame, pivot_df: pd.DataFrame, tipus: str, forras_nev: str) -> bytes:
    buf = io.BytesIO()
    summary_df = osszefoglalo_df(eredmeny_df)

    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        eredmeny_df.to_excel(writer, sheet_name="Összesítés", index=False, startrow=4)
        summary_df.to_excel(writer, sheet_name="Összefoglaló", index=False, startrow=4)
        pivot_df.to_excel(writer, sheet_name="Munkalaponként", index=False, startrow=4)
        laponkenti_df.to_excel(writer, sheet_name="Részletes", index=False, startrow=4)

        wb = writer.book
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            ws["A1"] = APP_TITLE
            ws["A2"] = f"Statisztika típusa: {tipus}"
            ws["A3"] = f"Forrás: {forras_nev}"
            ws["A1"].font = Font(bold=True, size=14)
            ws["A2"].font = Font(italic=True)
            ws["A3"].font = Font(italic=True)

            header_row = 5
            fill = PatternFill("solid", fgColor="D9EAF7")
            border = Border(bottom=Side(style="thin", color="808080"))
            for cell in ws[header_row]:
                cell.font = Font(bold=True)
                cell.fill = fill
                cell.border = border
                cell.alignment = Alignment(horizontal="center")

            for column_cells in ws.columns:
                max_len = 0
                col_letter = get_column_letter(column_cells[0].column)
                for cell in column_cells:
                    if cell.value is not None:
                        max_len = max(max_len, len(str(cell.value)))
                ws.column_dimensions[col_letter].width = min(max(max_len + 2, 12), 42)

            ws.freeze_panes = "A6"
            last_row = ws.max_row + 2
            ws[f"A{last_row}"] = FOOTER_TEXT
            ws[f"A{last_row}"].font = Font(italic=True, size=9)

    return buf.getvalue()


def main():
    st.set_page_config(page_title=APP_TITLE, layout="wide")
    st.title(APP_TITLE)
    st.caption("Streamlit felület Excel exporttal")

    uploaded = st.file_uploader("Excel fájl feltöltése", type=["xlsx"])
    tipus = st.radio("Statisztika típusa", ["féléves", "évvégi"], horizontal=True)

    if uploaded is None:
        st.info("Tölts fel egy Excel fájlt a statisztika elkészítéséhez.")
        st.stop()

    try:
        wb = load_workbook(io.BytesIO(uploaded.getvalue()), read_only=True, data_only=True)
        sheetnames = wb.sheetnames
        wb.close()
    except Exception as e:
        st.error(f"Nem sikerült beolvasni az Excel fájlt: {e}")
        st.stop()

    with st.expander("Munkalap szűrés", expanded=False):
        selected_sheets = st.multiselect("Mely munkalapokat vegye figyelembe?", sheetnames, default=sheetnames)

    if not selected_sheets:
        st.warning("Legalább egy munkalapot válassz ki.")
        st.stop()

    if st.button("Statisztika elkészítése", type="primary"):
        with st.spinner("Statisztika készítése..."):
            try:
                eredmeny_df, laponkenti_df, pivot_df = statisztika_keszitese(uploaded.getvalue(), tipus, selected_sheets)
                st.session_state["eredmeny_df"] = eredmeny_df
                st.session_state["laponkenti_df"] = laponkenti_df
                st.session_state["pivot_df"] = pivot_df
                st.session_state["tipus"] = tipus
                st.session_state["forras_nev"] = uploaded.name
            except Exception as e:
                st.error(f"Hiba a fájl feldolgozása közben: {e}")

    if "eredmeny_df" not in st.session_state:
        st.stop()

    eredmeny_df = st.session_state["eredmeny_df"]
    laponkenti_df = st.session_state["laponkenti_df"]
    pivot_df = st.session_state["pivot_df"]
    tipus = st.session_state["tipus"]
    forras_nev = st.session_state["forras_nev"]

    st.success("Statisztika elkészült.")

    col1, col2, col3 = st.columns(3)
    col1.metric("Összes bejegyzés", int(eredmeny_df["Darab"].sum()))
    col2.metric("Dicséretek", int(eredmeny_df.loc[eredmeny_df["Kategória"].isin(DICSERETEK), "Darab"].sum()))
    col3.metric("Fegyelmi / figyelmeztetés", int(eredmeny_df.loc[eredmeny_df["Kategória"].isin(FEGYELMI), "Darab"].sum()))

    st.subheader("Összesítés kategóriánként")
    st.dataframe(eredmeny_df, width="stretch", hide_index=True)

    chart_df = eredmeny_df[eredmeny_df["Darab"] > 0].copy()
    if not chart_df.empty:
        st.subheader("Diagram")
        st.bar_chart(chart_df.set_index("Kategória")["Darab"])

    with st.expander("Munkalaponkénti bontás"):
        st.dataframe(pivot_df, width="stretch", hide_index=True)

    xlsx_out = excel_export_bytes(eredmeny_df, laponkenti_df, pivot_df, tipus, forras_nev)
    st.download_button(
        "Excel export letöltése",
        data=xlsx_out,
        file_name=f"osztalystatisztika_{tipus}.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )

    st.caption(FOOTER_TEXT)


if __name__ == "__main__":
    main()
